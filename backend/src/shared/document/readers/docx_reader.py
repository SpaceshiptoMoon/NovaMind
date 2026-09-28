"""DOCX 文档读取器。"""
import os

from docx import Document
from novamind.shared.document.readers.base_reader import BaseReader
from novamind.shared.document.readers.executor import run_in_executor
from novamind.shared.logging import get_logger

logger = get_logger(__name__)


class DocxReader(BaseReader):
    """DOCX文档读取器"""

    def __init__(self):
        """初始化基类公共字段（编码列表、logger）。"""
        super().__init__()

    def _load_data_sync(self, file_path: str) -> list[dict[str, str]]:
        """用 python-docx 提取段落与表格文本并拼成整块。"""
        documents = []
        try:
            # 使用python-docx读取DOCX
            doc = Document(file_path)

            # 提取所有段落的文本
            full_text = []
            for para in doc.paragraphs:
                full_text.append(para.text)

            # 获取表格中的文本
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        full_text.append(cell.text)

            text = '\n'.join(full_text)

            if text.strip():
                documents.append({
                    'text': text,
                    'source': os.path.basename(file_path),
                    'page': 1,
                    'doc_id': f"docx_{hash(file_path)}",
                    'type': 'docx'
                })
        except Exception as e:
            logger.error("读取DOCX文件失败", file_path=str(file_path), error=str(e))

        return documents

    async def load_data(self, file_path: str) -> list[dict[str, str]]:
        """异步加载 DOCX 文本，重活转共享线程池避免阻塞事件循环。

        Args:
            file_path: DOCX 文件路径。

        Returns:
            单元素文档块列表（段落+表格全文拼合），读取失败时返回空列表。
        """
        return await run_in_executor(self._load_data_sync, file_path)