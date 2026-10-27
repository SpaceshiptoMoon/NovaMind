"""DeepDoc Excel 解析器。"""
from __future__ import annotations

import datetime
import logging
import re
from io import BytesIO
from pathlib import Path

import pandas as pd

# Adapted from RAGFlow deepdoc/parser/excel_parser.py

ILLEGAL_CHARACTERS_RE = re.compile(r"[\000-\010]|[\013-\014]|[\016-\037]")

logger = logging.getLogger(__name__)


class RAGFlowExcelParser:
    """RAGFlow Excel 解析器：工作表逐行转 Markdown 行文本。"""
    @staticmethod
    def _import_openpyxl():
        """懒加载 openpyxl。"""
        from openpyxl import Workbook, load_workbook

        return Workbook, load_workbook

    @staticmethod
    def _load_excel_to_workbook(file_like_object):
        """从路径/字节流加载为工作簿；非 Excel 载荷走 CSV 兜底。
        
        先按文件头魔数判断（zip 的 PK 头 / OLE 旧格式头），非 Excel 载荷用 pandas
        读 CSV；openpyxl 加载失败再降级 pandas（calamine 引擎兜底），统一包成工作簿。
        """
        _, load_workbook = RAGFlowExcelParser._import_openpyxl()
        # 路径形态先读入内存（真机实测：runtime_parser 传 Path 的 str 形态时
        # 无条件 seek(0) 崩溃——str 被当路径却从未打开），再统一包成字节流
        if isinstance(file_like_object, (str, Path)):
            file_like_object = Path(file_like_object).read_bytes()
        if isinstance(file_like_object, bytes):
            file_like_object = BytesIO(file_like_object)
        file_like_object.seek(0)
        file_head = file_like_object.read(4)
        file_like_object.seek(0)

        if not (file_head.startswith(b"PK\x03\x04") or file_head.startswith(b"\xd0\xcf\x11\xe0")):
            logging.info("DeepDoc excel parser received non-Excel payload; attempting CSV fallback")
            df = pd.read_csv(file_like_object, on_bad_lines="skip")
            return RAGFlowExcelParser._dataframe_to_workbook(df)

        try:
            wb = load_workbook(file_like_object, data_only=True)
        except Exception as exc:
            logging.info("openpyxl load failed, trying pandas fallback: %s", exc)
            file_like_object.seek(0)
            try:
                dfs = pd.read_excel(file_like_object, sheet_name=None)
            except Exception:
                file_like_object.seek(0)
                dfs = pd.read_excel(file_like_object, engine="calamine")
            return RAGFlowExcelParser._dataframe_to_workbook(dfs)
        RAGFlowExcelParser._expand_merged_cells(wb)
        RAGFlowExcelParser._backfill_formula_values(wb, file_like_object)
        return wb

    @staticmethod
    def _expand_merged_cells(wb):
        """合并单元格值展开：左上角值填充整个合并区域。

        openpyxl 里合并区域只有左上格有值、其余为 None——行遍历与 HTML 输出
        中合并维度会变成空串（如「华东区」合并 3 行，第 2/3 行区域列为空），
        行级语义断裂。展开后每行都携带区域值（代价是值在多行重复，属检索
        友好的冗余而非错误）。
        """
        for ws in wb.worksheets:
            merged_ranges = list(ws.merged_cells.ranges)
            if not merged_ranges:
                continue
            # 展开前快照：语义层（banner/小节行判定）需要原始合并信息
            setattr(
                ws,
                "_novamind_merged_pre",
                [(rng.min_row, rng.min_col, rng.max_row, rng.max_col) for rng in merged_ranges],
            )
            for rng in merged_ranges:
                top_left = ws.cell(row=rng.min_row, column=rng.min_col).value
                ws.unmerge_cells(str(rng))
                if top_left is None:
                    continue
                for row in range(rng.min_row, rng.max_row + 1):
                    for col in range(rng.min_col, rng.max_col + 1):
                        ws.cell(row=row, column=col).value = top_left

    @staticmethod
    def _backfill_formula_values(wb, file_like_object, size_limit_mb: int = 20):
        """公式无缓存值回填：data_only 视角为 None 且公式视图有值的格填公式字面量。

        程序生成的 xlsx（openpyxl/pandas 写出）只有公式没有缓存值，
        data_only=True 视角这些格为 None——检索产物丢值（常见于部分格是
        公式的常规报表，不限于整表公式）。以 data_only=False 二次加载，
        把公式字符串（如 =SUM(A1:A2)）当字面内容填回：诚实暴露公式而非
        静默丢值。文件超限不二次加载（内存防护）。
        """
        try:
            file_like_object.seek(0)
            head = file_like_object.read(4)
            file_like_object.seek(0)
        except Exception:
            return
        if not head.startswith(b"PK\x03\x04"):
            return  # CSV 兜底产物没有公式概念
        try:
            file_size = file_like_object.getbuffer().nbytes
        except AttributeError:
            return
        if file_size > size_limit_mb * 1024 * 1024:
            logging.info(
                "Excel formula backfill skipped: workbook too large (%d bytes)", file_size
            )
            return
        _, load_workbook = RAGFlowExcelParser._import_openpyxl()
        try:
            formula_wb = load_workbook(file_like_object, data_only=False)
        except Exception as exc:
            logging.info("Formula-view reload failed: %s", exc)
            return
        backfilled = 0
        for sheet_offset, ws in enumerate(wb.worksheets):
            formula_ws = formula_wb.worksheets[sheet_offset]
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        continue
                    formula_cell = formula_ws.cell(row=cell.row, column=cell.column)
                    if formula_cell.value is not None and str(formula_cell.value).strip():
                        cell.value = formula_cell.value
                        backfilled += 1
        if backfilled:
            logging.info("Excel formula backfill: %d cells", backfilled)

    @staticmethod
    def format_cell_value(value):
        """单元格值转展示文本：去浮点整尾、日期去零点时间、控制字符剔除。

        str(value) 的默认形态带检索噪音：1.0 显示「1.0」、日期显示
        「2024-01-01 00:00:00」——对 BM25 与阅读都是噪音。
        """
        if value is None:
            return ""
        if isinstance(value, bool):
            return "是" if value else "否"
        if isinstance(value, float) and value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        if isinstance(value, datetime.datetime) and (
            value.hour == 0 and value.minute == 0 and value.second == 0
        ):
            return value.date().isoformat()
        if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
            return str(value)
        if isinstance(value, str):
            return ILLEGAL_CHARACTERS_RE.sub(" ", value).strip()
        if isinstance(value, (list, tuple)):
            # openpyxl rich_text 模式的 CellRichText（list 子类）——str() 会把
            # 对象表示写进索引；拼接其文本片段（TextBlock 有 .text 属性）
            parts = [
                str(getattr(part, "text", None) if getattr(part, "text", None) else part)
                for part in value
            ]
            return ILLEGAL_CHARACTERS_RE.sub(" ", "".join(parts)).strip()
        return str(value)

    @staticmethod
    def _clean_dataframe(df: pd.DataFrame):
        """清理数据帧：剔除控制字符。"""
        def clean_string(value):
            """剔除单元格文本中的控制字符。"""
            if isinstance(value, str):
                return ILLEGAL_CHARACTERS_RE.sub(" ", value)
            return value

        return df.apply(lambda col: col.map(clean_string))

    @staticmethod
    def _fill_worksheet_from_dataframe(ws, df: pd.DataFrame):
        """把数据帧首行当表头、逐单元格写入工作表（pandas 兜底路径复用）。"""
        for col_num, column_name in enumerate(df.columns, 1):
            ws.cell(row=1, column=col_num, value=column_name)
        for row_num, row in enumerate(df.values, 2):
            for col_num, value in enumerate(row, 1):
                ws.cell(row=row_num, column=col_num, value=value)

    @staticmethod
    def _dataframe_to_workbook(df):
        """数据帧（或多表字典）包成工作簿，供统一的行遍历路径消费。"""
        Workbook, _ = RAGFlowExcelParser._import_openpyxl()
        if isinstance(df, dict) and len(df) > 1:
            return RAGFlowExcelParser._dataframes_to_workbook(df)

        df = RAGFlowExcelParser._clean_dataframe(df)
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        RAGFlowExcelParser._fill_worksheet_from_dataframe(ws, df)
        return wb

    @staticmethod
    def _dataframes_to_workbook(dfs: dict):
        """多数据帧合并为一个多表工作簿。"""
        Workbook, _ = RAGFlowExcelParser._import_openpyxl()
        wb = Workbook()
        default_sheet = wb.active
        wb.remove(default_sheet)

        for sheet_name, df in dfs.items():
            df = RAGFlowExcelParser._clean_dataframe(df)
            ws = wb.create_sheet(title=sheet_name)
            RAGFlowExcelParser._fill_worksheet_from_dataframe(ws, df)
        return wb

    @staticmethod
    def _get_actual_row_count(ws):
        """探测工作表实际数据行数。
        
        openpyxl 的 max_row 常含样式残留的虚增行；超大表先抽样确认有数据，
        再二分定位最后的数据行，尾段线性扫描收边。
        """
        max_row = ws.max_row
        if not max_row:
            return 0
        if max_row <= 10000:
            return max_row

        max_col = min(ws.max_column or 1, 50)

        def row_has_data(row_idx):
            """该行前 50 列是否存在非空单元格（虚增行探测）。"""
            for col_idx in range(1, max_col + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                if cell.value is not None and str(cell.value).strip():
                    return True
            return False

        if not any(row_has_data(i) for i in range(1, min(101, max_row + 1))):
            return 0

        left, right = 1, max_row
        last_data_row = 1
        while left <= right:
            mid = (left + right) // 2
            found = False
            for row_idx in range(mid, min(mid + 10, max_row + 1)):
                if row_has_data(row_idx):
                    found = True
                    last_data_row = max(last_data_row, row_idx)
                    break
            if found:
                left = mid + 1
            else:
                right = mid - 1

        for row_idx in range(last_data_row, min(last_data_row + 500, max_row + 1)):
            if row_has_data(row_idx):
                last_data_row = row_idx
        return last_data_row

    @staticmethod
    def _get_rows_limited(ws):
        """读取工作表实际数据范围内的行对象列表。"""
        actual_rows = RAGFlowExcelParser._get_actual_row_count(ws)
        if actual_rows == 0:
            return []
        return list(ws.iter_rows(min_row=1, max_row=actual_rows))

    _NUMERIC_LIKE_RE = re.compile(r"^-?[$￥€]?\d[\d,，]*\.?\d*%?$")

    @classmethod
    def _detect_header_rows(cls, rows) -> int:
        """探测表头行数（支持多级复合表头，最多向下看 3 行）。

        表头行的特征：非空格占比高且数值型格占比低（≤30%）；遇到数值主导行
        即停。首行就数值主导（占比 >70%）判为无表头表，返回 0——此时首行
        按数据行输出，表头用列字母占位，不再错把数据行当表头丢弃语义。
        """
        checked = 0
        header_rows = 0
        for row in list(rows)[:3]:
            values = [RAGFlowExcelParser.format_cell_value(c.value) for c in row]
            non_empty = [v for v in values if v]
            if len(non_empty) < 2:
                break
            numeric = [v for v in non_empty if cls._NUMERIC_LIKE_RE.match(v)]
            numeric_ratio = len(numeric) / len(non_empty)
            if checked == 0 and numeric_ratio > 0.7:
                return 0  # 无表头表
            if numeric_ratio > 0.3:
                break
            header_rows += 1
            checked += 1
        return header_rows or 1

    @classmethod
    def _build_composite_headers(cls, rows, header_rows: int) -> list[str]:
        """把前 header_rows 行拼成每列单一复合表头（「Q1 > 收入」形态）。

        合并单元格展开会让组头值在同行横向重复（Q1 填满其下所有列）——
        纵向拼接时跳过与前一段相同的值，得到「组头 > 子头」。全列无值时
        用列字母占位。
        """
        # 物化为 list：openpyxl 的 _Row 代理可迭代但无 len()
        header_block = [list(row) for row in rows[:header_rows]]
        if header_rows <= 0 or not header_block:
            first = list(rows[0]) if rows else []
            return [cls._column_letter(i) for i in range(len(first))]
        width = max(len(row) for row in header_block)
        headers: list[str] = []
        for col in range(width):
            pieces: list[str] = []
            for row in header_block:
                value = cls.format_cell_value(row[col].value) if col < len(row) else ""
                if value and value not in pieces:
                    pieces.append(value)
            headers.append(" > ".join(pieces) if pieces else cls._column_letter(col))
        return headers

    @staticmethod
    def _column_letter(index: int) -> str:
        """0-based 列序号转 Excel 列字母（0→A，25→Z，26→AA）。"""
        letters = ""
        index += 1
        while index:
            index, rem = divmod(index - 1, 26)
            letters = chr(ord("A") + rem) + letters
        return letters

    def html(self, fnm, chunk_rows=256):
        """工作表转 HTML 表格文本，按 chunk_rows 行分块（多级表头折成复合表头）。

        Args:
            fnm: 文件路径或字节流。
            chunk_rows: 每块包含的数据行数。

        Returns:
            HTML 表格字符串列表，每块带工作表名 caption。
        """
        blocks, _images = self.parse_blocks(fnm, chunk_rows=chunk_rows)
        return blocks

    def parse_blocks(self, fnm, chunk_rows=256):
        """工作表转 HTML 块 + 嵌入图片/图表提取（runtime 主入口，一次加载）。

        版面语义管线：值网格（excel_layout.build_grid）→ 多岛投影分割 →
        缝合后处理 → 逐岛前置区/表头识别 → 复合表头渲染（带 `Sheet!A1:F18`
        引用坐标 caption 与单元格批注内联）→ 嵌入图片占位符挂锚点所在岛、
        图表 range 回读缓存值重建数据表。

        Args:
            fnm: 文件路径或字节流。
            chunk_rows: 每块包含的数据行数（超行窗口重复表头）。

        Returns:
            (blocks, images) 二元组：blocks 为 HTML 块列表（表格/前置文本/
            图表摘要）；images 为嵌入图片字典列表（sheet/row/col/blob/
            artifact_id，png 化重编码，不可解码图跳过）。
        """
        from html import escape

        from novamind.engines.document.integrations.deepdoc.parsers import (
            excel_layout as layout,
        )
        _row_texts = layout._row_texts

        file_like_object = BytesIO(fnm) if not isinstance(fnm, str) else fnm
        wb = RAGFlowExcelParser._load_excel_to_workbook(file_like_object)
        tb_chunks: list[str] = []
        images: list[dict] = []

        def cell_html(text: str, note: str = "") -> str:
            """单元格转 HTML：换行转 <br/>，批注内联到格尾。"""
            body = escape(text).replace("\n", "<br/>")
            if note:
                body = f"{body} {escape(note)}"
            return f"<td>{body}</td>"

        # 阶段 1：全簿值网格（图表引用可能跨 sheet 回读）
        grids: dict[str, Any] = {}
        for sheetname in wb.sheetnames:
            try:
                grids[sheetname] = layout.build_grid(
                    wb[sheetname], RAGFlowExcelParser.format_cell_value
                )
            except Exception:
                logger.warning(
                    "Skip sheet '%s': grid build failed", sheetname, exc_info=True
                )

        for sheet_index, sheetname in enumerate(wb.sheetnames):
            grid = grids.get(sheetname)
            if grid is None:
                continue
            ws = wb[sheetname]
            sheet_images = RAGFlowExcelParser._extract_sheet_images(ws, sheetname)
            for order, image_info in enumerate(sheet_images):
                location = (
                    f"r{image_info['row'] + 1}c{image_info['col'] + 1}"
                    if image_info["row"] is not None and image_info["col"] is not None
                    else f"pos{order}"
                )
                image_info["artifact_id"] = f"xls_s{sheet_index}_{location}_{order}"

            islands = layout.stitch_islands(
                grid, layout.segment_islands(grid)
            )
            sheet_blocks: list[str] = []
            island_table_ranges: list[tuple[int, int, Island]] = []  # (起, 止, 岛)
            for island in islands:
                base = len(sheet_blocks)
                for line in island.preamble:
                    sheet_blocks.append(line)
                if island.bottom >= island.data_top:
                    # 单位语义：前置区「单位：万元」挂到 caption，数值检索有量纲
                    unit = layout.extract_unit(island.preamble)
                    unit_suffix = f"（单位：{escape(unit)}）" if unit else ""
                    transpose = layout.detect_transpose(grid, island, island.data_top)
                    sections = layout.section_rows(grid, island, island.data_top)
                    tl = f"{layout._column_letter(island.left)}{island.top + 1}"
                    br = f"{layout._column_letter(island.right)}{island.bottom + 1}"
                    caption = f"{sheetname}!{tl}:{br}{unit_suffix}"
                    if not island.has_header:
                        # 诚实标注：列字母是占位不是作者表头
                        caption += "（无表头，按列字母标注）"
                    if transpose:
                        caption += "（转置表：行头在首列）"
                    if transpose:
                        # 转置表：首列作行表头（th），列头用 band 或列字母
                        corner = island.headers[0] if island.has_header else ""
                        header_row = (
                            f"<tr><th>{escape(corner)}</th>"
                            + "".join(
                                f"<th>{escape(layout._column_letter(i))}</th>"
                                for i in range(island.right - island.left)
                            )
                            + "</tr>"
                        )
                    else:
                        header_row = (
                            "<tr>"
                            + "".join(f"<th>{escape(h)}</th>" for h in island.headers)
                            + "</tr>"
                        )
                    n_data = island.bottom - island.data_top + 1
                    table_open = False
                    table_html = ""
                    for window_i in range((n_data + chunk_rows - 1) // chunk_rows):
                        start = island.data_top + window_i * chunk_rows
                        for r in range(start, min(start + chunk_rows, island.bottom + 1)):
                            # 小节行：拆出为独立标注行（关闭当前表、开新表续排）
                            if r in sections:
                                if table_open:
                                    sheet_blocks.append(table_html + "</table>\n")
                                    table_open = False
                                sheet_blocks.append(f"[节] {escape(sections[r])}")
                                continue
                            if not table_open:
                                table_html = (
                                    f"<table><caption>{escape(caption)}</caption>{header_row}"
                                )
                                table_open = True
                            is_subtotal = layout.is_subtotal_row(
                                _row_texts(grid, r, island.left, island.right)
                            )
                            if transpose:
                                head = f"<th>{escape(grid.text_at(r, island.left))}</th>"
                                tds = "".join(
                                    cell_html(
                                        grid.text_at(r, c),
                                        grid.comments.get((r, c), ""),
                                    )
                                    for c in range(island.left + 1, island.right + 1)
                                )
                                table_html += f"<tr>{head}{tds}</tr>"
                            elif is_subtotal:
                                texts = _row_texts(grid, r, island.left, island.right)
                                cells = []
                                for offset, c in enumerate(range(island.left, island.right + 1)):
                                    text = grid.text_at(r, c)
                                    note = grid.comments.get((r, c), "")
                                    if offset == 0 and text:
                                        text = f"[小计] {text}"
                                    cells.append(cell_html(text, note))
                                table_html += f"<tr>{''.join(cells)}</tr>"
                            else:
                                tds = "".join(
                                    cell_html(
                                        grid.text_at(r, c),
                                        grid.comments.get((r, c), ""),
                                    )
                                    for c in range(island.left, island.right + 1)
                                )
                                table_html += f"<tr>{tds}</tr>"
                    if table_open:
                        sheet_blocks.append(table_html + "</table>\n")
                island_table_ranges.append((base, len(sheet_blocks), island))

            # 嵌入图片：占位符挂锚点所在岛的首块；岛外浮图挂锚点上方最近的岛
            # （表尾浮图通常注解该表）；上方无岛才独立成块
            for image_info in sheet_images:
                placeholder = f"__FIGURE_URL__{image_info['artifact_id']}__\n"
                r0, c0 = image_info.get("row"), image_info.get("col")
                placed = False
                if r0 is not None and c0 is not None:
                    for base, end, island in island_table_ranges:
                        if (
                            island.top <= r0 <= island.bottom
                            and island.left <= c0 <= island.right
                            and end > base
                        ):
                            sheet_blocks[base] = placeholder + sheet_blocks[base]
                            placed = True
                            break
                if not placed and r0 is not None:
                    above = [
                        (base, end)
                        for base, end, island in island_table_ranges
                        if island.top <= r0 and end > base
                    ]
                    if above:
                        base, _end = above[-1]
                        sheet_blocks[base] = placeholder + sheet_blocks[base]
                        placed = True
                if not placed:
                    sheet_blocks.append(placeholder.strip())
            images.extend(sheet_images)

            # 图表：range 引用回读缓存值重建数据表
            for chart in getattr(ws, "_charts", []) or []:
                block = self._chart_to_block(chart, sheetname, grids)
                if block:
                    sheet_blocks.append(block)

            tb_chunks.extend(sheet_blocks)
        return tb_chunks, images

    def _chart_to_block(self, chart, sheetname: str, grids: dict) -> str:
        """图表对象 → 数据表 HTML：series range 引用回读缓存值重建
        「类别列 + 系列列」表（对齐 docling `_chart_to_table_data` 的思路，
        零外部依赖）。重建失败/无数值返回空串（诚实跳过）。
        """
        from html import escape
        from io import BytesIO

        from openpyxl.utils.cell import range_boundaries
        from PIL import Image as PILImage

        ctype = type(chart).__name__.replace("Chart", "").lower() or "chart"
        title = RAGFlowExcelParser._chart_title_text(chart)

        def resolve_values(ref) -> list[str]:
            """把 Reference/numRef/strRef 的 range 引用解析为网格展示文本。"""
            formula = getattr(ref, "f", None) or (
                getattr(getattr(ref, "numRef", None), "f", None)
            )
            if not formula:
                return []
            sheet_part, _, range_part = str(formula).rpartition("!")
            sheet_name = sheet_part.strip("'")
            grid = grids.get(sheet_name)
            if grid is None:
                return []
            try:
                min_col, min_row, max_col, max_row = range_boundaries(range_part)
            except Exception:
                return []
            if (max_row - min_row + 1) * (max_col - min_col + 1) > 100_000:
                return []
            out: list[str] = []
            for r in range(min_row - 1, max_row):
                for c in range(min_col - 1, max_col):
                    out.append(grid.text_at(r, c))
            return out

        series_rows: list[tuple[str, list[str]]] = []
        categories: list[str] = []
        max_len = 0
        for ser in getattr(chart, "series", []) or []:
            name_ref = getattr(getattr(ser, "tx", None), "strRef", None)
            name_values = resolve_values(name_ref)
            name = name_values[0] if name_values else "系列"
            val_ref = getattr(getattr(ser, "val", None), "numRef", None)
            values = resolve_values(val_ref)
            if not values:
                continue
            cat = getattr(ser, "cat", None)
            cat_values = resolve_values(
                getattr(cat, "strRef", None) or getattr(cat, "numRef", None)
            ) if cat is not None else []
            if cat_values and not categories:
                categories = cat_values
            series_rows.append((name, values))
            max_len = max(max_len, len(values))
        if not series_rows or max_len == 0:
            return ""
        if max_len * len(series_rows) > 100_000:
            logger.warning("Chart data too large to rebuild, skipped (sheet=%s)", sheetname)
            return ""

        rows_html = ""
        for i in range(max_len):
            cat = categories[i] if i < len(categories) else ""
            tds = f"<td>{escape(cat)}</td>" + "".join(
                f"<td>{escape(vals[i] if i < len(vals) else '')}</td>"
                for _name, vals in series_rows
            )
            rows_html += f"<tr>{tds}</tr>"
        header = "".join(f"<th>{escape(name)}</th>" for name, _ in series_rows)
        caption = f"图表（{escape(ctype)}）：{escape(title or sheetname)}（{sheetname}）"
        return (
            f"<table><caption>{caption}</caption>"
            f"<tr><th>类别</th>{header}</tr>{rows_html}</table>\n"
        )

    @staticmethod
    def _chart_title_text(chart) -> str:
        """从图表 Title 富文本对象提取标题文本（取不到返回空串）。"""
        try:
            parts: list[str] = []
            for paragraph in chart.title.tx.rich.p:
                for run in getattr(paragraph, "r", None) or []:
                    if run.t:
                        parts.append(run.t)
            return "".join(parts).strip()
        except Exception:
            return ""

    @staticmethod
    def _extract_sheet_images(ws, sheetname: str) -> list[dict]:
        """提取工作表嵌入图片：png 化重编码（统一 figure 管线格式）。

        EMF/WMF 等 Windows 元文件 Pillow 不可解码，跳过并告警（诚实暴露
        而非产出坏图）。
        """
        from io import BytesIO

        from PIL import Image as PILImage

        out: list[dict] = []
        for order, img in enumerate(getattr(ws, "_images", []) or []):
            try:
                data = img._data()
            except Exception:
                logger.warning(
                    "Skip unreadable embedded image in sheet '%s' (order %d)",
                    sheetname, order, exc_info=True,
                )
                continue
            if not data:
                continue
            row = col = None
            anchor = getattr(img, "anchor", None)
            marker = getattr(anchor, "_from", None)
            if marker is not None:
                row = int(getattr(marker, "row", 0))
                col = int(getattr(marker, "col", 0))
            try:
                pil = PILImage.open(BytesIO(data))
                pil.load()
                if pil.mode not in ("RGB", "RGBA", "L"):
                    pil = pil.convert("RGBA")
                buf = BytesIO()
                pil.save(buf, format="PNG")
                png_bytes = buf.getvalue()
            except Exception:
                logger.warning(
                    "Skip undecodable embedded image in sheet '%s' (order %d, "
                    "likely EMF/WMF)", sheetname, order, exc_info=True,
                )
                continue
            out.append({
                "sheet": sheetname,
                "row": row,
                "col": col,
                "blob": png_bytes,
                # artifact_id 由 runtime_parser 组装（需 sheet 序号，此处不含）
            })
        return out

    def __call__(self, fnm):
        """解析 Excel 为「复合表头：值；……」行文本列表（sheet 名作后缀）。

        与 html 出口同源的多级表头处理：前 N 表头行纵向折成复合表头，
        数据行从表头块之后开始；无表头表用列字母做前缀。
        """
        file_like_object = BytesIO(fnm) if not isinstance(fnm, str) else fnm
        wb = RAGFlowExcelParser._load_excel_to_workbook(file_like_object)

        rows_out = []
        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            try:
                rows = RAGFlowExcelParser._get_rows_limited(ws)
            except Exception:
                logger.warning("Skip sheet '%s' due to rows access error", sheetname, exc_info=True)
                continue
            if not rows:
                continue
            header_rows = RAGFlowExcelParser._detect_header_rows(rows)
            headers = RAGFlowExcelParser._build_composite_headers(rows, header_rows)
            for row in list(rows[header_rows:]):
                fields = []
                for idx, cell in enumerate(row):
                    value_text = RAGFlowExcelParser.format_cell_value(cell.value)
                    if not value_text:
                        continue
                    title = headers[idx] if idx < len(headers) else ""
                    prefix = f"{title}：" if title else ""
                    fields.append(prefix + value_text)
                if not fields:
                    continue
                line = "; ".join(fields)
                if "sheet" not in sheetname.lower():
                    line += f" —— {sheetname}"
                rows_out.append(line)
        return rows_out
