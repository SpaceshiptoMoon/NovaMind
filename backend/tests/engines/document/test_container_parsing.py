"""容器递归解析回归：zip 压缩包与 eml/msg 邮件附件的展开、防御与合并。

2026-10 容器解析批次：企业上传常见「zip 里一堆文档」「邮件带附件」——
解析层展开成员逐个递归（成员天然过各自入口守卫），产物按成员标题头拼
接、figure region 重编号并入容器盐。防御：非 zip 拒收 / 成员数与解压
总量超限拒收 / 嵌套 zip 深度限制 / 打包元数据过滤 / 成员失败降级占位。
"""
import base64
import email
import email.policy
import sys
import zipfile
from io import BytesIO

import pytest

BACKEND_ROOT = Path = __import__("pathlib").Path(__file__).resolve().parents[2]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

pytestmark = pytest.mark.unit

_PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
)


def _parser():
    from novamind.engines.document.integrations.deepdoc.core.runtime_parser import (
        DeepDocParser,
    )

    return DeepDocParser()


def _parse_bytes(payload: bytes, file_type: str):
    import asyncio

    return asyncio.run(_parser().parse_bytes(payload, file_type=file_type))


def _make_docx(text: str, *, with_image: bool = False) -> bytes:
    """构造最小 docx 字节流（可选内嵌 1px 图片触发 figure 契约）。"""
    from docx import Document
    from docx.shared import Inches

    doc = Document()
    doc.add_paragraph(text)
    if with_image:
        doc.add_picture(BytesIO(_PNG_1PX), width=Inches(1))
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_zip(entries: dict[str, bytes]) -> bytes:
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, blob in entries.items():
            zf.writestr(name, blob)
    return buf.getvalue()


# ---------- zip 正常路径 ----------

def test_zip_members_parsed_and_joined():
    """zip 内 txt+docx：成员各自解析，标题头进正文，members 清单出账。"""
    payload = _make_zip({
        "说明.txt": "压缩包说明文字".encode("utf-8"),
        "文档.docx": _make_docx("内嵌文档正文"),
    })
    result = _parse_bytes(payload, "zip")
    assert result.metadata["members"] == ["说明.txt", "文档.docx"]
    assert "【附件】说明.txt" in result.full_text
    assert "压缩包说明文字" in result.full_text
    assert "【附件】文档.docx" in result.full_text
    assert "内嵌文档正文" in result.full_text
    assert result.chunks  # 容器级统一切分出块


def test_zip_figure_regions_renumbered_and_resalted():
    """两个带图 docx：region 重编号 m0_/m1_ 前缀，正文占位符换容器盐版本。"""
    import re

    payload = _make_zip({
        "a.docx": _make_docx("第一份", with_image=True),
        "b.docx": _make_docx("第二份", with_image=True),
    })
    result = _parse_bytes(payload, "zip")
    regions = result.metadata["figure_regions"]
    assert [r["artifact_id"] for r in regions] == ["m0_docx_fig_0", "m1_docx_fig_0"]
    # 容器盐一致性：全部 region 的 placeholder 同盐（同一 12hex 前缀）
    salts = {re.search(r"__FIGURE_URL__([0-9a-f]{12})_", r["placeholder"]).group(1) for r in regions}
    assert len(salts) == 1
    # 正文中的占位符与 region 同构（成员旧占位符已被精确替换，无残留）
    for region in regions:
        assert region["placeholder"] in result.full_text
    assert not re.search(r"__FIGURE_URL__[0-9a-f]{12}_docx_fig_0__", result.full_text.replace(
        regions[0]["placeholder"], "").replace(regions[1]["placeholder"], ""))


# ---------- zip 防御 ----------

def test_zip_not_zip_rejected():
    """非 zip 内容拒收（PK 魔数守卫）。"""
    from novamind.engines.document.exceptions import MismatchedPayloadError

    with pytest.raises(MismatchedPayloadError):
        _parse_bytes(b"<html>fake zip</html>", "zip")


def test_zip_member_failure_degrades():
    """坏 docx 成员解析失败：降级占位说明，其余成员不受牵连。"""
    payload = _make_zip({
        "broken.docx": b"<html>not a real docx</html>",
        "ok.txt": "正常成员内容".encode("utf-8"),
    })
    result = _parse_bytes(payload, "zip")
    assert "正常成员内容" in result.full_text
    assert "【附件】broken.docx" in result.full_text
    assert "未能解析" in result.full_text


def test_zip_noise_entries_filtered():
    """__MACOSX/.DS_Store/隐藏文件不入成员流。"""
    payload = _make_zip({
        "__MACOSX/._a.txt": b"junk",
        ".DS_Store": b"junk",
        ".hidden.txt": b"junk",
        "real.txt": "真实内容".encode("utf-8"),
    })
    result = _parse_bytes(payload, "zip")
    assert result.metadata["members"] == ["real.txt"]


def test_zip_nested_one_level_expanded():
    """zip 套 zip：深度上限内递归展开，内层成员直出。"""
    inner = _make_zip({"inner.txt": "内层成员".encode("utf-8")})
    payload = _make_zip({"outer.zip": inner, "outer.txt": "外层成员".encode("utf-8")})
    result = _parse_bytes(payload, "zip")
    assert "inner.txt" in result.metadata["members"]
    assert "内层成员" in result.full_text


def test_zip_member_count_limit_rejected(monkeypatch):
    """成员数超限拒收整包（不静默截断）。"""
    from novamind.engines.document.exceptions import MismatchedPayloadError
    from novamind.engines.document.pipeline import nested_archive

    monkeypatch.setattr(nested_archive, "ARCHIVE_MAX_MEMBERS", 2)
    payload = _make_zip({f"f{i}.txt": f"内容{i}".encode("utf-8") for i in range(3)})
    with pytest.raises(MismatchedPayloadError, match="成员数超过上限"):
        _parse_bytes(payload, "zip")


# ---------- eml 附件递归 ----------

def _make_eml(attachments: dict[str, bytes], body: str = "邮件正文内容") -> bytes:
    msg = email.message.EmailMessage()
    msg["From"] = "sender@example.com"
    msg["To"] = "receiver@example.com"
    msg["Subject"] = "测试邮件"
    msg.set_content(body)
    for name, blob in attachments.items():
        msg.add_attachment(blob, maintype="application", subtype="octet-stream", filename=name)
    return msg.as_bytes()


def test_eml_attachment_recursively_parsed():
    """eml 带 txt 附件：正文与附件内容都进产物，附件成员清单出账。"""
    payload = _make_eml({"报告.txt": "附件里的报告内容".encode("utf-8")})
    result = _parse_bytes(payload, "eml")
    assert "邮件正文内容" in result.full_text
    assert "【附件】报告.txt" in result.full_text
    assert "附件里的报告内容" in result.full_text
    assert result.metadata["members"] == ["报告.txt"]


def test_eml_attachment_limit_rejected(monkeypatch):
    """eml 附件数超限走同一限额拒收。"""
    from novamind.engines.document.exceptions import MismatchedPayloadError
    from novamind.engines.document.pipeline import nested_archive

    monkeypatch.setattr(nested_archive, "ARCHIVE_MAX_MEMBERS", 1)
    payload = _make_eml({f"a{i}.txt": f"内容{i}".encode("utf-8") for i in range(2)})
    with pytest.raises(MismatchedPayloadError, match="成员数超过上限"):
        _parse_bytes(payload, "eml")
