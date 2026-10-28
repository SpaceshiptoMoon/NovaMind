"""多格式图片提取回归：pptx/epub/mineru 的 figure 契约全覆盖。

2026-10-28 图片提取全覆盖批次：此前仅 pdf/docx/xlsx 产出 figure_regions，
pptx（PICTURE 形状整图丢失）、epub（章节内嵌图丢失）、mineru（产物目录
图丢失且引用成死链）三个缺口补齐——统一契约（带盐占位符 + image_blobs）
交宿主上传链。
"""
import base64
import sys
import tempfile
import zipfile
from io import BytesIO
from pathlib import Path

import pytest

BACKEND_ROOT = Path(__file__).resolve().parents[2]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

pytestmark = pytest.mark.unit

_PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
)

_SALT_RE = None  # 占位符形态断言用（惰性编译见下）


def _salt_re():
    import re

    global _SALT_RE
    if _SALT_RE is None:
        _SALT_RE = re.compile(r"__FIGURE_URL__[0-9a-f]{12}_")
    return _SALT_RE


# ---------- pptx ----------


def test_pptx_picture_shape_extracted():
    """PICTURE 形状：blob 进 figures（带页码），占位符进页文本流。"""
    from pptx import Presentation
    from pptx.util import Inches

    from novamind.engines.document.integrations.deepdoc.parsers.ppt import (
        RAGFlowPptParser,
    )

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.add_picture(BytesIO(_PNG_1PX), Inches(1), Inches(1))
    path = Path(tempfile.mkdtemp()) / "pic.pptx"
    prs.save(path)

    slides, figures = RAGFlowPptParser()(str(path))
    assert len(figures) == 1
    assert figures[0]["artifact_id"] == "ppt_fig_0"
    assert figures[0]["pages"] == [1] and figures[0]["page_start"] == 1
    assert figures[0]["image_blobs"] == [_PNG_1PX]
    assert _salt_re().search(slides[0]), "占位符应在页文本流"
    assert f"]({figures[0]['placeholder']})" in slides[0]


# ---------- epub ----------


def _build_epub_with_image() -> Path:
    path = Path(tempfile.mkdtemp()) / "book.epub"
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("mimetype", "application/epub+zip")
        zf.writestr(
            "META-INF/container.xml",
            '<?xml version="1.0"?><container xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
            '<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>'
            "</rootfiles></container>",
        )
        zf.writestr(
            "OEBPS/content.opf",
            '''<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="2.0">
<manifest><item id="c1" href="ch1.xhtml" media-type="application/xhtml+xml"/>
<item id="c2" href="ch2.xhtml" media-type="application/xhtml+xml"/>
<item id="img1" href="images/pic.png" media-type="image/png"/></manifest>
<spine><itemref idref="c1"/><itemref idref="c2"/></spine></package>''',
        )
        zf.writestr(
            "OEBPS/ch1.xhtml",
            '<html><body><p>章节正文段落。</p><img src="images/pic.png"/></body></html>',
        )
        zf.writestr("OEBPS/ch2.xhtml", "<html><body><p>无图章节。</p></body></html>")
        zf.writestr("OEBPS/images/pic.png", _PNG_1PX)
    return path


def test_epub_chapter_image_extracted():
    """章节 <img>：blob 经 manifest 映射提取，占位符附所属章节文本尾。"""
    from novamind.engines.document.integrations.deepdoc.parsers.epub import (
        RAGFlowEpubParser,
    )

    sections, figures = RAGFlowEpubParser()(str(_build_epub_with_image()))
    assert len(figures) == 1
    assert figures[0]["artifact_id"] == "epub_fig_0"
    assert figures[0]["image_blobs"] == [_PNG_1PX]
    joined = "\n".join(sections)
    assert "章节正文段落" in joined
    assert f"]({figures[0]['placeholder']})" in joined


def test_epub_img_src_traversal_rejected():
    """img src 路径穿越（未在 manifest image 集合内的 zip 路径）不提取。"""
    from novamind.engines.document.integrations.deepdoc.parsers.epub import (
        RAGFlowEpubParser,
    )

    path = Path(tempfile.mkdtemp()) / "evil.epub"
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("mimetype", "application/epub+zip")
        zf.writestr(
            "META-INF/container.xml",
            '<?xml version="1.0"?><container xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
            '<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>'
            "</rootfiles></container>",
        )
        zf.writestr(
            "OEBPS/content.opf",
            '''<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="2.0">
<manifest><item id="c1" href="ch1.xhtml" media-type="application/xhtml+xml"/></manifest>
<spine><itemref idref="c1"/></spine></package>''',
        )
        zf.writestr(
            "OEBPS/ch1.xhtml",
            '<html><body><img src="../../etc/secret.png"/><p>正文</p></body></html>',
        )
        zf.writestr("secret.png", _PNG_1PX)  # 包外的敏感路径模拟

    sections, figures = RAGFlowEpubParser()(str(path))
    assert figures == [], "穿越路径不应被提取"
    assert any("正文" in s for s in sections)


# ---------- mineru ----------


def test_mineru_figure_contract_shape():
    """mineru 契约：figures 字段默认空元组（dataclass 兼容既有构造）。"""
    from novamind.engines.document.integrations.mineru.client import MineruParseResult

    result = MineruParseResult(full_text="x", page_count=1)
    assert result.figures == ()
    result2 = MineruParseResult(
        full_text="x",
        page_count=1,
        figures=(
            {
                "artifact_id": "mineru_fig_0",
                "placeholder": "__FIGURE_URL__a1b2c3d4e5f6_mineru_fig_0__",
                "image_blobs": [_PNG_1PX],
            },
        ),
    )
    assert len(result2.figures) == 1


def test_mineru_markdown_image_link_rebound_to_placeholder():
    """markdown 相对图片引用就地替换为带盐占位符（收集逻辑同款串操作）。"""
    from novamind.engines.document.integrations.deepdoc.core.document_model import (
        document_salt,
        figure_placeholder,
    )

    full_text = "正文 ![x](images/a.jpg) 结尾"
    salt = document_salt(b"pdf-bytes")
    placeholder = figure_placeholder(salt, "mineru_fig_0")
    replaced = full_text.replace("](images/a.jpg)", f"]({placeholder})")
    assert placeholder in replaced
    assert "images/a.jpg" not in replaced


# ---------- xlsx（端到端：此前 placeholder 未透传导致替换永不命中） ----------


def test_xlsx_figure_region_carries_placeholder_and_e2e_replacement():
    """xlsx region 必须携带带盐 placeholder，业务替换/presign 正则全链命中。

    回归：盐化批次只改了解析器侧（正文占位符带盐、image_info.placeholder），
    runtime 组装 region 时未透传 placeholder——业务 map key 兜底无盐形态与
    正文带盐占位符永不命中，图片上传成功却进不了 MD/chunk（用户实测发现）。
    """
    import asyncio
    import re

    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XImage

    from novamind.engines.document.integrations.deepdoc.core.runtime_parser import (
        DeepDocParser,
    )
    from novamind.features.knowledge_space.services.document_pipeline import (
        _restore_figure_links,
    )

    png_path = Path(tempfile.mkdtemp()) / "a.png"
    png_path.write_bytes(_PNG_1PX)
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "数据表"
    img = XImage(png_path)
    img.anchor = "B5"
    ws.add_image(img)
    xlsx = Path(tempfile.mkdtemp()) / "e.xlsx"
    wb.save(xlsx)

    result = asyncio.run(
        DeepDocParser().parse(xlsx, parsing_config={}, splitting_config={"chunk_size": 1000})
    )
    regions = result.metadata.get("figure_regions") or []
    assert regions and regions[0].get("placeholder"), "region 必须带 placeholder"
    assert regions[0]["placeholder"] in result.full_text, "正文应含带盐占位符"

    url_map = {regions[0]["placeholder"]: "figure_x_0.png"}
    valid = {r["placeholder"] for r in regions}
    replaced, _ = _restore_figure_links(
        result.full_text, list(result.chunks), url_map, valid
    )
    assert "figure_x_0.png" in replaced, "业务替换应命中带盐占位符"
    assert not re.search(r"__FIGURE_URL__[0-9a-f]{12}_", replaced), "占位符应剥净"
    assert re.findall(r"\]\((figure_[A-Za-z0-9_.-]+\.png)\)", replaced) == [
        "figure_x_0.png"
    ], "替换后应匹配读取端点 presign 正则"


# ---------- docx 纯图片段 + 契约形状统一 ----------


def test_docx_standalone_image_paragraph_extracted():
    """纯图片段（图片独占一段，Word 常见形态）：此前 text 判定先行整段消失。"""
    import asyncio
    from docx import Document as DocxDocument

    from novamind.engines.document.integrations.deepdoc.core.runtime_parser import (
        DeepDocParser,
    )

    doc = DocxDocument()
    doc.add_paragraph("带图文档")
    doc.add_paragraph().add_run().add_picture(BytesIO(_PNG_1PX))
    path = Path(tempfile.mkdtemp()) / "solo.docx"
    doc.save(path)

    result = asyncio.run(
        DeepDocParser().parse(path, parsing_config={}, splitting_config={"chunk_size": 500})
    )
    regions = result.metadata.get("figure_regions") or []
    assert len(regions) == 1, "纯图片段应产出 figure region"
    assert result.metadata["images"] == 1
    assert _salt_re().search(result.full_text)


def test_contract_shape_uniform_across_formats(tmp_path):
    """不管什么文件，metadata 形状一致：figure_regions/document_blocks 恒存在，
    region 必需键齐全（下游写一种逻辑，无需逐格式探测）。"""
    import asyncio
    from docx import Document as DocxDocument

    from novamind.engines.document.integrations.deepdoc.core.runtime_parser import (
        DeepDocParser,
    )

    (tmp_path / "a.txt").write_text("纯文本内容，无图。", encoding="utf-8")
    doc = DocxDocument()
    doc.add_paragraph("带图文档")
    doc.add_paragraph().add_run().add_picture(BytesIO(_PNG_1PX))
    doc.save(tmp_path / "b.docx")

    parser = DeepDocParser()
    results = {
        "txt": asyncio.run(
            parser.parse(tmp_path / "a.txt", parsing_config={}, splitting_config={"chunk_size": 500})
        ),
        "docx": asyncio.run(
            parser.parse(tmp_path / "b.docx", parsing_config={}, splitting_config={"chunk_size": 500})
        ),
    }
    required = (
        "artifact_id", "placeholder", "pages", "page_start", "bbox",
        "caption", "text", "member_texts", "member_bboxes", "image_blobs",
    )
    for fmt, result in results.items():
        metadata = result.metadata
        assert "figure_regions" in metadata and "document_blocks" in metadata, (
            f"{fmt} 缺契约键"
        )
        for region in metadata["figure_regions"]:
            missing = [key for key in required if key not in region]
            assert not missing, f"{fmt} region 缺键: {missing}"
            assert region["placeholder"], f"{fmt} placeholder 为空"
    # docx 有图、txt 无图（空数组兜底）
    assert len(results["docx"].metadata["figure_regions"]) == 1
    assert results["txt"].metadata["figure_regions"] == []
