"""解析质量报告回归：WARNING 捕获、并发隔离、metadata 出账。"""
import asyncio
import logging
import sys
import threading
from pathlib import Path

import pytest

BACKEND_ROOT = Path(__file__).resolve().parents[2]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

pytestmark = pytest.mark.unit

from novamind.engines.document.integrations.deepdoc.core.quality import (
    build_quality_metadata,
    collect_warnings,
)
from novamind.engines.document.integrations.deepdoc.core.runtime_parser import (
    DeepDocParser,
)


def test_collect_warnings_captures_thread_warnings():
    """当前线程的 WARNING+ 被捕获，INFO 与截断语义正确。"""
    with collect_warnings() as warnings:
        logging.getLogger("test.q").info("info 不收")
        logging.getLogger("test.q").warning("第一条警告 %s", "带参数")
        logging.getLogger("test.q").error("x" * 500)
    assert len(warnings) == 2
    assert warnings[0] == "WARNING: 第一条警告 带参数"
    assert warnings[1].startswith("ERROR: ") and warnings[1].endswith("...")
    assert len(warnings[1]) <= 307  # 消息体 300 + 级别前缀 7


def test_collect_warnings_isolated_across_threads():
    """并发解析各自线程的警告互不串话（worker max_jobs>1 场景）。"""
    results: dict[int, list[str]] = {}
    barrier = threading.Barrier(2)

    def _worker(tag: str) -> None:
        with collect_warnings() as warnings:
            barrier.wait()  # 两个 collector 同时在挂载状态
            logging.getLogger("test.q").warning(f"来自{tag}")
        results[threading.get_ident()] = warnings

    threads = [threading.Thread(target=_worker, args=(t,)) for t in ("甲", "乙")]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert sorted(w[0] for w in results.values()) == [
        "WARNING: 来自乙",
        "WARNING: 来自甲",
    ]


def test_build_quality_metadata_shape():
    """quality 块契约：parser/warnings/duration_ms + 特有字段。"""
    import time

    quality = build_quality_metadata("deepdoc", ["w1"], time.perf_counter() - 0.05, pdf_mode="full")
    assert quality["parser"] == "deepdoc"
    assert quality["warnings"] == ["w1"]
    assert quality["duration_ms"] >= 40
    assert quality["pdf_mode"] == "full"


def _build_simple_docx(path: Path) -> None:
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_heading("质量报告夹具", level=1)
    doc.add_paragraph("正文段落。")
    doc.save(path)


def test_parse_result_metadata_carries_quality(tmp_path):
    """runtime 解析产物的 metadata 带 quality 块（parser/warnings/duration）。"""
    docx = tmp_path / "q.docx"
    _build_simple_docx(docx)
    result = asyncio.run(
        DeepDocParser().parse(docx, parsing_config={}, splitting_config={"chunk_size": 500})
    )
    quality = result.metadata["quality"]
    assert quality["parser"] == "deepdoc"
    assert isinstance(quality["warnings"], list)
    assert quality["duration_ms"] >= 0
